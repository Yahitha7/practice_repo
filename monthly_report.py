import os
from datetime import datetime

import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)
from openpyxl.utils import get_column_letter
from openpyxl.chart import (
    BarChart,
    LineChart,
    Reference
)
from openpyxl.chart.label import DataLabelList


# ============================================================
# PATH CONFIGURATION
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

HISTORY_DIR = os.path.join(
    BASE_DIR,
    "history"
)

DAILY_HISTORY_FILE = os.path.join(
    HISTORY_DIR,
    "daily_capacity.csv"
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "output"
)


# ============================================================
# MONTHLY CONFIGURATION
# ============================================================

NUMBER_OF_MONTHS = 3

# Backup team thresholds
TH_LOW = 0.80
TH_MEDIUM = 0.90
TH_HIGH = 0.95


# ============================================================
# EXCEL STYLES
# ============================================================

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

HEADER_FONT = Font(
    bold=True,
    color="000000"
)

TITLE_FONT = Font(
    bold=True,
    size=16
)

THIN_SIDE = Side(
    style="thin",
    color="D9E1F2"
)

BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE
)

CENTER = Alignment(
    horizontal="center",
    vertical="center"
)


# ============================================================
# LOAD HISTORY
# ============================================================

def load_history():

    if not os.path.exists(
        DAILY_HISTORY_FILE
    ):
        raise FileNotFoundError(
            "Historical file not found: "
            + DAILY_HISTORY_FILE
        )

    print(
        "Historical file used : "
        + DAILY_HISTORY_FILE
    )

    df = pd.read_csv(
        DAILY_HISTORY_FILE,
        keep_default_na=False
    )

    if df.empty:
        raise ValueError(
            "Historical file is empty: "
            + DAILY_HISTORY_FILE
        )

    required_columns = [
        "Date",
        "Cluster",
        "Active",
        "Location",
        "Type",
        "Region",
        "Country",
        "City",
        "Total Capacity (GB)",
        "Used Capacity (GB)",
        "Free Capacity (GB)",
        "Used Capacity (%)",
        "Free Capacity (%)"
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in df.columns
    ]

    if missing_columns:
        raise ValueError(
            "Historical file is missing columns: "
            + ", ".join(missing_columns)
        )

    # --------------------------------------------------------
    # Date
    # --------------------------------------------------------

    df["Date"] = pd.to_datetime(
        df["Date"],
        errors="coerce"
    )

    df = df.dropna(
        subset=["Date"]
    )

    # --------------------------------------------------------
    # Numeric columns
    # --------------------------------------------------------

    numeric_columns = [
        "Total Capacity (GB)",
        "Used Capacity (GB)",
        "Free Capacity (GB)",
        "Used Capacity (%)",
        "Free Capacity (%)"
    ]

    for column in numeric_columns:

        df[column] = pd.to_numeric(
            df[column],
            errors="coerce"
        )

    # --------------------------------------------------------
    # Normalize text
    # --------------------------------------------------------

    text_columns = [
        "Cluster",
        "Active",
        "Location",
        "Type",
        "Region",
        "Country",
        "City"
    ]

    for column in text_columns:

        df[column] = (
            df[column]
            .astype(str)
            .str.strip()
        )

    # --------------------------------------------------------
    # Restore Region from cluster mapping
    # --------------------------------------------------------

    try:

        from cluster_mapping import CLUSTER_MAPPING

        for index in df.index:

            cluster = df.at[
                index,
                "Cluster"
            ]

            region = df.at[
                index,
                "Region"
            ]

            if (
                region == ""
                or str(region).lower() == "nan"
            ):

                mapping = CLUSTER_MAPPING.get(
                    cluster
                )

                if mapping:

                    df.at[
                        index,
                        "Region"
                    ] = mapping.get(
                        "Region",
                        ""
                    )

                    if (
                        df.at[
                            index,
                            "Country"
                        ] == ""
                    ):

                        df.at[
                            index,
                            "Country"
                        ] = mapping.get(
                            "Country",
                            ""
                        )

                    if (
                        df.at[
                            index,
                            "City"
                        ] == ""
                    ):

                        df.at[
                            index,
                            "City"
                        ] = mapping.get(
                            "City",
                            ""
                        )

    except ImportError:

        pass

    # --------------------------------------------------------
    # Remove invalid clusters
    # --------------------------------------------------------

    df = df[
        (df["Cluster"] != "")
        &
        (df["Cluster"].str.lower() != "nan")
    ]

    # --------------------------------------------------------
    # Sort
    # --------------------------------------------------------

    df = (
        df
        .sort_values(
            by=[
                "Date",
                "Cluster"
            ]
        )
        .reset_index(
            drop=True
        )
    )

    print(
        "Historical rows loaded : "
        + str(len(df))
    )

    print(
        "Unique dates           : "
        + str(df["Date"].nunique())
    )

    print(
        "Unique clusters        : "
        + str(df["Cluster"].nunique())
    )

    return df


# ============================================================
# GET MONTHS
# ============================================================

def get_latest_months(df):

    if df.empty:
        return []

    data = df.copy()

    data["Month_Period"] = (
        data["Date"]
        .dt.to_period("M")
    )

    available_months = sorted(
        data["Month_Period"].unique()
    )

    selected_months = (
        available_months[
            -NUMBER_OF_MONTHS:
        ]
    )

    return selected_months


# ============================================================
# BUILD MONTHLY TABLE
# ============================================================

def build_monthly_table(
    history,
    selected_months
):

    if history.empty:
        return pd.DataFrame()

    monthly_records = []

    for month_period in selected_months:

        month_data = history[
            history["Date"].dt.to_period("M")
            == month_period
        ].copy()

        if month_data.empty:
            continue

        # Latest available snapshot in the month
        # for every cluster.

        month_data = (
            month_data
            .sort_values("Date")
            .groupby(
                "Cluster",
                as_index=False
            )
            .tail(1)
        )

        month_data = month_data[
            [
                "Cluster",
                "Region",
                "Country",
                "City",
                "Used Capacity (%)"
            ]
        ].copy()

        month_data["Month_Period"] = (
            month_period
        )

        monthly_records.append(
            month_data
        )

    if not monthly_records:
        return pd.DataFrame()

    monthly_data = pd.concat(
        monthly_records,
        ignore_index=True
    )

    # --------------------------------------------------------
    # Pivot months
    # --------------------------------------------------------

    pivot = monthly_data.pivot_table(
        index=[
            "Cluster",
            "Region",
            "Country",
            "City"
        ],
        columns="Month_Period",
        values="Used Capacity (%)",
        aggfunc="last"
    ).reset_index()

    pivot.columns.name = None

    # --------------------------------------------------------
    # Rename months
    # --------------------------------------------------------

    rename_map = {}

    for month_period in selected_months:

        if month_period in pivot.columns:

            rename_map[
                month_period
            ] = pd.Timestamp(
                month_period.start_time
            ).strftime(
                "%b-%y"
            )

    pivot = pivot.rename(
        columns=rename_map
    )

    # --------------------------------------------------------
    # Month columns
    # --------------------------------------------------------

    month_columns = []

    for month_period in selected_months:

        month_name = pd.Timestamp(
            month_period.start_time
        ).strftime(
            "%b-%y"
        )

        if month_name in pivot.columns:

            month_columns.append(
                month_name
            )

    # --------------------------------------------------------
    # Final columns
    # --------------------------------------------------------

    result_columns = [
        "Cluster",
        "Region",
        "Country",
        "City"
    ]

    result_columns.extend(
        month_columns
    )

    pivot = pivot[
        result_columns
    ]

    # --------------------------------------------------------
    # Sort
    # --------------------------------------------------------

    pivot = (
        pivot
        .sort_values(
            by=[
                "City",
                "Cluster"
            ]
        )
        .reset_index(
            drop=True
        )
    )

    return pivot


# ============================================================
# CREATE REGIONAL SHEET
# ============================================================

def create_region_sheet(
    wb,
    region_data,
    region_name,
    month_columns
):

    ws = wb.create_sheet(
        title=region_name
    )

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    ws["A1"] = (
        "Rubrik Monthly Capacity Growth"
    )

    ws["A1"].font = TITLE_FONT

    # --------------------------------------------------------
    # No Filters
    # --------------------------------------------------------

    # --------------------------------------------------------
    # Main table
    # --------------------------------------------------------

    start_row = 5

    headers = [
        "City Location",
        "TH %-80",
        "TH %-90",
        "TH %-95"
    ]

    headers.extend(
        month_columns
    )

    for column_number, header in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row=start_row,
            column=column_number,
            value=header
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER

    # --------------------------------------------------------
    # Right-side chart source table
    # --------------------------------------------------------

    right_start_col = 10

    for column_offset, header in enumerate(
        headers,
        start=right_start_col
    ):

        cell = ws.cell(
            row=start_row,
            column=column_offset,
            value=header
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER

    # --------------------------------------------------------
    # Write regional data
    # --------------------------------------------------------

    table_data = region_data.copy()

    if not table_data.empty:

        table_data["City Location"] = (
            table_data["City"]
            .astype(str)
            .str.strip()
        )

        table_data = table_data[
            table_data["City Location"] != ""
        ]

        table_data = table_data.sort_values(
            by=[
                "City Location",
                "Cluster"
            ]
        )

    for row_offset, (_, record) in enumerate(
        table_data.iterrows(),
        start=1
    ):

        excel_row = (
            start_row
            + row_offset
        )

        values = [
            record["City Location"],
            TH_LOW,
            TH_MEDIUM,
            TH_HIGH
        ]

        for month in month_columns:

            values.append(
                record.get(
                    month,
                    None
                )
            )

        # Main table
        for column_number, value in enumerate(
            values,
            start=1
        ):

            cell = ws.cell(
                row=excel_row,
                column=column_number,
                value=value
            )

            cell.border = BORDER

            if column_number >= 2:

                cell.number_format = "0%"

        # Chart source table
        for offset, value in enumerate(
            values,
            start=right_start_col
        ):

            cell = ws.cell(
                row=excel_row,
                column=offset,
                value=value
            )

            cell.border = BORDER

            if offset >= right_start_col + 1:

                cell.number_format = "0%"

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14

    for column_number in range(
        4,
        4 + len(month_columns)
    ):

        ws.column_dimensions[
            get_column_letter(column_number)
        ].width = 14

    ws.column_dimensions["J"].width = 22

    for column_number in range(
        right_start_col + 1,
        right_start_col + 4 + len(month_columns)
    ):

        ws.column_dimensions[
            get_column_letter(column_number)
        ].width = 14

    # --------------------------------------------------------
    # Chart
    # --------------------------------------------------------

    if not table_data.empty:

        first_data_row = start_row + 1

        last_data_row = (
            start_row
            + len(table_data)
        )

        # ----------------------------------------------------
        # Bar chart
        # ----------------------------------------------------

        bar_chart = BarChart()

        bar_chart.type = "col"

        bar_chart.title = ""

        bar_chart.y_axis.numFmt = "0%"

        bar_chart.y_axis.scaling.min = 0
        bar_chart.y_axis.scaling.max = 1
        bar_chart.y_axis.majorUnit = 0.25

        bar_chart.x_axis.title = ""

        categories = Reference(
            ws,
            min_col=right_start_col,
            min_row=first_data_row,
            max_row=last_data_row
        )

        # ----------------------------------------------------
        # Add one bar series for every month
        # ----------------------------------------------------

        for index, month in enumerate(
            month_columns
        ):

            month_column = (
                right_start_col
                + 4
                + index
            )

            values = Reference(
                ws,
                min_col=month_column,
                min_row=start_row,
                max_row=last_data_row
            )

            bar_chart.add_data(
                values,
                titles_from_data=True
            )

        bar_chart.set_categories(
            categories
        )

        bar_chart.dLbls = DataLabelList()

        bar_chart.dLbls.showVal = True

        bar_chart.dLbls.numFmt = "0%"

        bar_chart.dLbls.position = "outEnd"

        # ----------------------------------------------------
        # Threshold lines
        # ----------------------------------------------------

        low_line = LineChart()

        low_values = Reference(
            ws,
            min_col=right_start_col + 1,
            min_row=start_row,
            max_row=last_data_row
        )

        low_line.add_data(
            low_values,
            titles_from_data=True
        )

        low_line.set_categories(
            categories
        )

        medium_line = LineChart()

        medium_values = Reference(
            ws,
            min_col=right_start_col + 2,
            min_row=start_row,
            max_row=last_data_row
        )

        medium_line.add_data(
            medium_values,
            titles_from_data=True
        )

        medium_line.set_categories(
            categories
        )

        high_line = LineChart()

        high_values = Reference(
            ws,
            min_col=right_start_col + 3,
            min_row=start_row,
            max_row=last_data_row
        )

        high_line.add_data(
            high_values,
            titles_from_data=True
        )

        high_line.set_categories(
            categories
        )

        # ----------------------------------------------------
        # Combine
        # ----------------------------------------------------

        bar_chart += low_line
        bar_chart += medium_line
        bar_chart += high_line

        bar_chart.height = 10
        bar_chart.width = 18

        bar_chart.legend.position = "t"

        # ----------------------------------------------------
        # Insert chart below table
        # ----------------------------------------------------

        chart_row = (
            start_row
            + len(table_data)
            + 3
        )

        ws.add_chart(
            bar_chart,
            f"E{chart_row}"
        )

    ws.freeze_panes = "A6"

    return ws


# ============================================================
# CREATE RAW DATA SHEET
# ============================================================

def create_raw_data_sheet(
    wb,
    history
):

    ws = wb.create_sheet(
        title="Raw Data"
    )

    ws["A1"] = (
        "Rubrik Monthly Capacity Raw History"
    )

    ws["A1"].font = TITLE_FONT

    # --------------------------------------------------------
    # Raw data columns
    # --------------------------------------------------------

    columns = [
        "Date",
        "Month",
        "Cluster",
        "Active",
        "Location",
        "Type",
        "Region",
        "Country",
        "City",
        "Total Capacity (GB)",
        "Used Capacity (GB)",
        "Free Capacity (GB)",
        "Used Capacity (%)",
        "Free Capacity (%)",
        "TH High",
        "TH Medium",
        "TH Low"
    ]

    start_row = 3

    for column_number, header in enumerate(
        columns,
        start=1
    ):

        cell = ws.cell(
            row=start_row,
            column=column_number,
            value=header
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER

    # --------------------------------------------------------
    # Write history
    # --------------------------------------------------------

    sorted_history = history.sort_values(
        by=[
            "Date",
            "Cluster"
        ]
    ).reset_index(
        drop=True
    )

    for row_offset, (_, record) in enumerate(
        sorted_history.iterrows(),
        start=1
    ):

        row_number = (
            start_row
            + row_offset
        )

        values = [
            record["Date"],
            record["Date"].strftime("%Y-%m"),
            record["Cluster"],
            record["Active"],
            record["Location"],
            record["Type"],
            record["Region"],
            record["Country"],
            record["City"],
            record["Total Capacity (GB)"],
            record["Used Capacity (GB)"],
            record["Free Capacity (GB)"],
            record["Used Capacity (%)"],
            record["Free Capacity (%)"],
            TH_HIGH,
            TH_MEDIUM,
            TH_LOW
        ]

        for column_number, value in enumerate(
            values,
            start=1
        ):

            cell = ws.cell(
                row=row_number,
                column=column_number,
                value=value
            )

            cell.border = BORDER

            if column_number == 1:

                cell.number_format = "yyyy-mm-dd"

            elif column_number in [
                10,
                11,
                12
            ]:

                cell.number_format = "#,##0.00"

            elif column_number in [
                13,
                14,
                15,
                16,
                17
            ]:

                cell.number_format = "0%"

    # --------------------------------------------------------
    # Widths
    # --------------------------------------------------------

    widths = {
        "A": 14,
        "B": 12,
        "C": 16,
        "D": 12,
        "E": 40,
        "F": 16,
        "G": 12,
        "H": 18,
        "I": 20,
        "J": 20,
        "K": 20,
        "L": 20,
        "M": 18,
        "N": 18,
        "O": 12,
        "P": 12,
        "Q": 12
    }

    for column, width in widths.items():

        ws.column_dimensions[
            column
        ].width = width

    ws.freeze_panes = "A4"

    return ws


# ============================================================
# GENERATE MONTHLY REPORT
# ============================================================

def generate_monthly_report():

    print()
    print("=" * 70)
    print("Rubrik Monthly Capacity Report")
    print("=" * 70)

    # --------------------------------------------------------
    # Load history
    # --------------------------------------------------------

    history = load_history()

    # --------------------------------------------------------
    # Get latest months
    # --------------------------------------------------------

    selected_months = get_latest_months(
        history
    )

    print()

    print(
        "Months included : "
        + str(len(selected_months))
    )

    for month in selected_months:

        print(
            "  "
            + pd.Timestamp(
                month.start_time
            ).strftime(
                "%B %Y"
            )
        )

    # --------------------------------------------------------
    # Build monthly table
    # --------------------------------------------------------

    monthly_table = build_monthly_table(
        history,
        selected_months
    )

    print()

    print(
        "Clusters included : "
        + str(len(monthly_table))
    )

    if monthly_table.empty:

        raise ValueError(
            "No monthly capacity data available."
        )

    # --------------------------------------------------------
    # Month columns
    # --------------------------------------------------------

    month_columns = []

    for month in selected_months:

        month_name = pd.Timestamp(
            month.start_time
        ).strftime(
            "%b-%y"
        )

        if month_name in monthly_table.columns:

            month_columns.append(
                month_name
            )

    print(
        "Monthly columns    : "
        + ", ".join(
            month_columns
        )
    )

    # --------------------------------------------------------
    # Create output directory
    # --------------------------------------------------------

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    # --------------------------------------------------------
    # File name
    # --------------------------------------------------------

    today = datetime.now()

    report_name = today.strftime(
        "Backup_Capacity_Monthly_%d_%b_%Y.xlsx"
    )

    report_path = os.path.join(
        OUTPUT_DIR,
        report_name
    )

    # --------------------------------------------------------
    # Workbook
    # --------------------------------------------------------

    wb = Workbook()

    default_sheet = wb.active

    wb.remove(
        default_sheet
    )

    # --------------------------------------------------------
    # Regional sheets
    # --------------------------------------------------------

    regions = [
        "NA",
        "NA_DD",
        "LATAM",
        "EU",
        "APAC",
        "CHINA"
    ]

    for region in regions:

        if region == "CHINA":

            region_data = monthly_table[
                monthly_table["Country"]
                .astype(str)
                .str.upper()
                .eq("CHINA")
            ].copy()

        elif region == "NA_DD":

            region_data = monthly_table.iloc[
                0:0
            ].copy()

        else:

            region_data = monthly_table[
                monthly_table["Region"]
                .astype(str)
                .str.upper()
                .eq(region)
            ].copy()

        print(
            f"{region:<8}: "
            f"{len(region_data)} clusters"
        )

        create_region_sheet(
            wb,
            region_data,
            region,
            month_columns
        )

    # --------------------------------------------------------
    # Raw Data
    # --------------------------------------------------------

    create_raw_data_sheet(
        wb,
        history
    )

    # --------------------------------------------------------
    # Sheet order
    # --------------------------------------------------------

    desired_order = [
        "NA",
        "NA_DD",
        "LATAM",
        "EU",
        "APAC",
        "CHINA",
        "Raw Data"
    ]

    wb._sheets = [
        wb[name]
        for name in desired_order
        if name in wb.sheetnames
    ]

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    wb.save(
        report_path
    )

    file_exists = os.path.exists(
        report_path
    )

    print()
    print("=" * 70)
    print("Monthly Excel Report Generated")
    print("=" * 70)

    print(
        "Report Location : "
        + report_path
    )

    print(
        "Absolute Path   : "
        + os.path.abspath(
            report_path
        )
    )

    print(
        "Clusters        : "
        + str(len(monthly_table))
    )

    print(
        "Months Included : "
        + str(len(selected_months))
    )

    print(
        "File Exists     : "
        + str(file_exists)
    )

    print(
        "Sheets          : "
        + ", ".join(
            wb.sheetnames
        )
    )

    print("=" * 70)

    return report_path


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    report_file = (
        generate_monthly_report()
    )

    print()

    print(
        "Sending Monthly Capacity Report Email..."
    )

    from email_report import send_email

    send_email(
        report_file
    )
