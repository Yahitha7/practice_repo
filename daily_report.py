import os
from datetime import datetime

import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo
)
from openpyxl.chart import (
    BarChart,
    LineChart,
    Reference
)
from openpyxl.chart.label import DataLabelList


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "output"
)

# Backup team's thresholds
TH_LOW = 0.80
TH_MEDIUM = 0.90
TH_HIGH = 0.95


# ============================================================
# EXCEL STYLES
# ============================================================

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

HEADER_FONT = Font(
    bold=True,
    color="000000"
)

TITLE_FONT = Font(
    bold=True,
    size=14
)

THIN_SIDE = Side(
    style="thin",
    color="D9E1F2"
)

BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE
)

CENTER = Alignment(
    horizontal="center",
    vertical="center"
)


# ============================================================
# PREPARE DATA
# ============================================================

def prepare_dataframe(df):

    data = df.copy()

    # --------------------------------------------------------
    # Normalize column names
    # --------------------------------------------------------

    rename_map = {
        "Total Capacity":
            "Total Capacity (GB)",

        "Used Capacity":
            "Used Capacity (GB)",

        "Free Capacity":
            "Free Capacity (GB)",

        "Used %":
            "Used Capacity (%)",

        "Free %":
            "Free Capacity (%)"
    }

    data = data.rename(
        columns=rename_map
    )


    # --------------------------------------------------------
    # Ensure Cluster column is populated
    # --------------------------------------------------------

    if "Cluster" not in data.columns:
        data["Cluster"] = ""

    if "Target Storage" in data.columns:
        data["Cluster"] = data["Cluster"].where(
            data["Cluster"].astype(str).str.strip() != "",
            data["Target Storage"]
        )


    # --------------------------------------------------------
    # Actual report date
    # --------------------------------------------------------

    report_date = datetime.now().strftime(
        "%Y-%m-%d"
    )

    data["Date"] = report_date

    # Keep Month for Raw Data because backup
    # team's source data contains it.

    data["Month"] = datetime.now().strftime(
        "%Y-%m"
    )

    # --------------------------------------------------------
    # Backup team thresholds
    # --------------------------------------------------------

    data["TH High"] = TH_HIGH
    data["TH Medium"] = TH_MEDIUM
    data["TH Low"] = TH_LOW

    # --------------------------------------------------------
    # Make sure required columns exist
    # --------------------------------------------------------

    required_columns = [
        "Date",
        "Month",
        "Active",
        "Location",
        "Type",
        "Cluster",
        "Region",
        "Country",
        "City",
        "Total Capacity (GB)",
        "Used Capacity (GB)",
        "Free Capacity (GB)",
        "Used Capacity (%)",
        "Free Capacity (%)",
        "TH High",
        "TH Medium",
        "TH Low"
    ]

    for column in required_columns:

        if column not in data.columns:

            data[column] = ""

    return data


# ============================================================
# WRITE CELL
# ============================================================

def write_cell(
    ws,
    row,
    column,
    value,
    header=False,
    percent=False
):

    cell = ws.cell(
        row=row,
        column=column,
        value=value
    )

    cell.border = BORDER

    if header:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    elif percent:

        cell.number_format = "0%"

    return cell


# ============================================================
# CREATE REGIONAL SHEET
# ============================================================

def create_region_sheet(
    wb,
    region_data,
    region_name,
    report_date
):

    ws = wb.create_sheet(
        title=region_name
    )

    # --------------------------------------------------------
    # If no data
    # --------------------------------------------------------

    if region_data.empty:

        ws["A1"] = (
            f"No data available for {region_name}"
        )

        return ws

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    ws["A1"] = (
        "Rubrik Daily Capacity Growth"
    )

    ws["A1"].font = TITLE_FONT

    # --------------------------------------------------------
    # No Filters
    # --------------------------------------------------------


    # --------------------------------------------------------
    # Main table
    # --------------------------------------------------------

    start_row = 5

    headers = [
        "City Location",
        "TH %-80",
        "TH %-90",
        "TH %-95",
        report_date
    ]

    for column_number, header in enumerate(
        headers,
        start=1
    ):

        write_cell(
            ws,
            start_row,
            column_number,
            header,
            header=True
        )

    # --------------------------------------------------------
    # Prepare city data
    # --------------------------------------------------------

    table_data = region_data.copy()

    table_data["City Location"] = (
        table_data["City"]
        .astype(str)
        .str.strip()
    )

    table_data = table_data[
        table_data["City Location"]
        != ""
    ]

    table_data = table_data.sort_values(
        by="City Location"
    )

    # --------------------------------------------------------
    # Write regional rows
    # --------------------------------------------------------

    for row_offset, (_, row) in enumerate(
        table_data.iterrows(),
        start=1
    ):

        excel_row = (
            start_row
            + row_offset
        )

        city = row[
            "City Location"
        ]

        used_percentage = row[
            "Used Capacity (%)"
        ]

        write_cell(
            ws,
            excel_row,
            1,
            city
        )

        write_cell(
            ws,
            excel_row,
            2,
            TH_LOW,
            percent=True
        )

        write_cell(
            ws,
            excel_row,
            3,
            TH_MEDIUM,
            percent=True
        )

        write_cell(
            ws,
            excel_row,
            4,
            TH_HIGH,
            percent=True
        )

        write_cell(
            ws,
            excel_row,
            5,
            used_percentage,
            percent=True
        )

    # --------------------------------------------------------
    # Right-side duplicate table
    #
    # Backup team's report has the same data table
    # on the right side for chart source.
    # --------------------------------------------------------

    right_start_col = 10

    for column_offset, header in enumerate(
        headers,
        start=right_start_col
    ):

        write_cell(
            ws,
            start_row,
            column_offset,
            header,
            header=True
        )

    for row_offset, (_, row) in enumerate(
        table_data.iterrows(),
        start=1
    ):

        excel_row = (
            start_row
            + row_offset
        )

        city = row[
            "City Location"
        ]

        used_percentage = row[
            "Used Capacity (%)"
        ]

        write_cell(
            ws,
            excel_row,
            right_start_col,
            city
        )

        write_cell(
            ws,
            excel_row,
            right_start_col + 1,
            TH_LOW,
            percent=True
        )

        write_cell(
            ws,
            excel_row,
            right_start_col + 2,
            TH_MEDIUM,
            percent=True
        )

        write_cell(
            ws,
            excel_row,
            right_start_col + 3,
            TH_HIGH,
            percent=True
        )

        write_cell(
            ws,
            excel_row,
            right_start_col + 4,
            used_percentage,
            percent=True
        )

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14

    ws.column_dimensions["J"].width = 22
    ws.column_dimensions["K"].width = 14
    ws.column_dimensions["L"].width = 14
    ws.column_dimensions["M"].width = 14
    ws.column_dimensions["N"].width = 14

    # --------------------------------------------------------
    # Create chart
    # --------------------------------------------------------

    if not table_data.empty:

        first_data_row = start_row + 1

        last_data_row = (
            start_row
            + len(table_data)
        )

        # ----------------------------------------------------
        # Bar chart
        # ----------------------------------------------------

        bar_chart = BarChart()

        bar_chart.type = "col"

        bar_chart.style = 10

        bar_chart.title = ""

        bar_chart.y_axis.title = ""

        bar_chart.x_axis.title = ""

        bar_chart.y_axis.numFmt = "0%"

        bar_chart.y_axis.scaling.min = 0
        bar_chart.y_axis.scaling.max = 1
        bar_chart.y_axis.majorUnit = 0.25

        # Current date values

        current_values = Reference(
            ws,
            min_col=right_start_col + 4,
            min_row=start_row,
            max_row=last_data_row
        )

        categories = Reference(
            ws,
            min_col=right_start_col,
            min_row=first_data_row,
            max_row=last_data_row
        )

        bar_chart.add_data(
            current_values,
            titles_from_data=True
        )

        bar_chart.set_categories(
            categories
        )

        # Data labels

        if bar_chart.series:

            bar_chart.series[0].graphicalProperties.noFill = False

            bar_chart.series[0].dLbls = DataLabelList()

            bar_chart.series[0].dLbls.showVal = True
            bar_chart.series[0].dLbls.numFmt = "0%"
            bar_chart.series[0].dLbls.position = "outEnd"

        # ----------------------------------------------------
        # Low threshold line
        # ----------------------------------------------------

        low_line = LineChart()

        low_values = Reference(
            ws,
            min_col=right_start_col + 1,
            min_row=start_row,
            max_row=last_data_row
        )

        low_line.add_data(
            low_values,
            titles_from_data=True
        )

        low_line.set_categories(
            categories
        )

        # ----------------------------------------------------
        # Medium threshold line
        # ----------------------------------------------------

        medium_line = LineChart()

        medium_values = Reference(
            ws,
            min_col=right_start_col + 2,
            min_row=start_row,
            max_row=last_data_row
        )

        medium_line.add_data(
            medium_values,
            titles_from_data=True
        )

        medium_line.set_categories(
            categories
        )

        # ----------------------------------------------------
        # High threshold line
        # ----------------------------------------------------

        high_line = LineChart()

        high_values = Reference(
            ws,
            min_col=right_start_col + 3,
            min_row=start_row,
            max_row=last_data_row
        )

        high_line.add_data(
            high_values,
            titles_from_data=True
        )

        high_line.set_categories(
            categories
        )

        # ----------------------------------------------------
        # Combine charts
        # ----------------------------------------------------

        bar_chart += low_line
        bar_chart += medium_line
        bar_chart += high_line

        bar_chart.height = 10
        bar_chart.width = 18

        # ----------------------------------------------------
        # Legend
        # ----------------------------------------------------

        bar_chart.legend.position = "t"

        # ----------------------------------------------------
        # Insert chart
        # ----------------------------------------------------

        chart_row = (
            start_row
            + len(table_data)
            + 3
        )

        ws.add_chart(
            bar_chart,
            f"E{chart_row}"
        )

    ws.freeze_panes = "A6"

    return ws


# ============================================================
# CREATE RAW DATA SHEET
# ============================================================

def create_raw_data_sheet(
    wb,
    data
):

    ws = wb.create_sheet(
        "Raw Data"
    )

    # --------------------------------------------------------
    # Raw Data columns
    #
    # Keep Date explicitly.
    # --------------------------------------------------------

    raw_columns = [
        "Date",
        "Month",
        "Active",
        "Location",
        "Type",
        "Cluster",
        "Region",
        "Country",
        "City",
        "Total Capacity (GB)",
        "Used Capacity (GB)",
        "Free Capacity (GB)",
        "Used Capacity (%)",
        "Free Capacity (%)",
        "TH High",
        "TH Medium",
        "TH Low"
    ]

    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    for column_number, header in enumerate(
        raw_columns,
        start=1
    ):

        write_cell(
            ws,
            1,
            column_number,
            header,
            header=True
        )

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    for row_number, (_, row) in enumerate(
        data.iterrows(),
        start=2
    ):

        for column_number, column in enumerate(
            raw_columns,
            start=1
        ):

            value = row.get(
                column,
                ""
            )

            is_percent = (
                "Capacity (%)" in column
                or column in [
                    "TH High",
                    "TH Medium",
                    "TH Low"
                ]
            )

            write_cell(
                ws,
                row_number,
                column_number,
                value,
                percent=is_percent
            )

    # --------------------------------------------------------
    # Widths
    # --------------------------------------------------------

    widths = {
        "A": 14,
        "B": 12,
        "C": 12,
        "D": 40,
        "E": 16,
        "F": 18,
        "G": 12,
        "H": 16,
        "I": 20,
        "J": 20,
        "K": 20,
        "L": 20,
        "M": 18,
        "N": 18,
        "O": 12,
        "P": 12,
        "Q": 12,
        "R": 12
    }

    for column, width in widths.items():

        ws.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------------
    # Excel table
    # --------------------------------------------------------

    if ws.max_row >= 2:

        table = Table(
            displayName="TblRawData",
            ref=ws.dimensions
        )

        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        table.tableStyleInfo = style

        ws.add_table(
            table
        )

    ws.freeze_panes = "A2"

    return ws


# ============================================================
# GENERATE DAILY REPORT
# ============================================================

def generate_daily_report(
    df,
    history_df=None
):

    print()
    print("=" * 70)
    print("Generating Daily Excel Report")
    print("=" * 70)

    # --------------------------------------------------------
    # Output directory
    # --------------------------------------------------------

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Prepare current snapshot
    # --------------------------------------------------------

    data = prepare_dataframe(
        df
    )

    report_date = data.iloc[0]["Date"]

    print(
        f"Report Date : {report_date}"
    )

    print(
        f"Clusters    : {len(data)}"
    )

    # --------------------------------------------------------
    # Create workbook
    # --------------------------------------------------------

    wb = Workbook()

    default_sheet = wb.active

    wb.remove(
        default_sheet
    )

    # --------------------------------------------------------
    # Regional sheets
    # --------------------------------------------------------

    regions = [
        "NA",
        "LATAM",
        "EU",
        "APAC",
        "CHINA"
    ]

    for region in regions:

        if region == "CHINA":

            region_data = data[
                data["Country"]
                .astype(str)
                .str.upper()
                .eq("CHINA")
            ].copy()

        elif region == "NA_DD":

            region_data = data.iloc[
                0:0
            ].copy()

        else:

            region_data = data[
                data["Region"]
                .astype(str)
                .str.upper()
                .eq(region)
            ].copy()

        print(
            f"{region:<8}: "
            f"{len(region_data)} clusters"
        )

        create_region_sheet(
            wb,
            region_data,
            region,
            report_date
        )

    # --------------------------------------------------------
    # Raw Data
    # --------------------------------------------------------

    create_raw_data_sheet(
        wb,
        data
    )

    # --------------------------------------------------------
    # Sheet order
    #
    # NO GLOBAL VIEW
    # --------------------------------------------------------

    desired_order = [
        "NA",
        "LATAM",
        "EU",
        "APAC",
        "CHINA",
        "Raw Data"
    ]

    wb._sheets = [
        wb[name]
        for name in desired_order
        if name in wb.sheetnames
    ]

    # --------------------------------------------------------
    # File name
    # --------------------------------------------------------

    today = datetime.now()

    file_name = today.strftime(
        "Backup_Capacity_Daily_%d_%b_%Y.xlsx"
    )

    report_path = os.path.join(
        OUTPUT_DIR,
        file_name
    )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    wb.save(
        report_path
    )

    # --------------------------------------------------------
    # Verification
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("Daily Excel Report Generated")
    print("=" * 70)

    print(
        f"Report Location : {report_path}"
    )

    print(
        f"Clusters        : {len(data)}"
    )

    print(
        "Sheets          : "
        + ", ".join(
            wb.sheetnames
        )
    )

    print(
        f"File Exists     : "
        f"{os.path.exists(report_path)}"
    )

    print("=" * 70)

    return report_path


# ============================================================
# DIRECT EXECUTION
# ============================================================

if __name__ == "__main__":

    print(
        "daily_report.py loaded successfully."
    )
