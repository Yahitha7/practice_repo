import os
from datetime import datetime

import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint


# ============================================================
# PATH CONFIGURATION
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

HISTORY_FILE = os.path.join(
    BASE_DIR,
    "history",
    "daily_capacity.csv"
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "output"
)


# ============================================================
# WEEKLY CONFIGURATION
# ============================================================

NUMBER_OF_WEEKS = 4

TH_LOW = 0.30
TH_HIGH = 0.85


# ============================================================
# EXCEL STYLES
# ============================================================

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

HEADER_FONT = Font(
    bold=True
)

TITLE_FONT = Font(
    bold=True,
    size=14
)

NORMAL_FONT = Font(
    size=10
)

THIN_SIDE = Side(
    style="thin",
    color="D9E1F2"
)

BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE
)

CENTER = Alignment(
    horizontal="center",
    vertical="center"
)

LEFT = Alignment(
    horizontal="left",
    vertical="center"
)


# ============================================================
# LOAD HISTORY
# ============================================================

def load_history():

    if not os.path.exists(HISTORY_FILE):

        raise FileNotFoundError(
            "Historical file not found: "
            + HISTORY_FILE
        )

    print()
    print(
        "Historical file used : "
        + HISTORY_FILE
    )

    df = pd.read_csv(
        HISTORY_FILE,
        keep_default_na=False
    )

    if df.empty:

        raise ValueError(
            "Historical file is empty: "
            + HISTORY_FILE
        )

    required_columns = [
        "Date",
        "Cluster",
        "Active",
        "Location",
        "Type",
        "Region",
        "Country",
        "City",
        "Total Capacity (GB)",
        "Used Capacity (GB)",
        "Free Capacity (GB)",
        "Used Capacity (%)",
        "Free Capacity (%)"
    ]

    missing = [
        column
        for column in required_columns
        if column not in df.columns
    ]

    if missing:

        raise ValueError(
            "Historical file is missing columns: "
            + ", ".join(missing)
        )

    # --------------------------------------------------------
    # Date
    # --------------------------------------------------------

    df["Date"] = pd.to_datetime(
        df["Date"],
        errors="coerce"
    )

    df = df.dropna(
        subset=["Date"]
    )

    # --------------------------------------------------------
    # String fields
    # --------------------------------------------------------

    string_columns = [
        "Cluster",
        "Active",
        "Location",
        "Type",
        "Region",
        "Country",
        "City"
    ]

    for column in string_columns:

        df[column] = (
            df[column]
            .astype(str)
            .str.strip()
        )

    # --------------------------------------------------------
    # Numeric fields
    # --------------------------------------------------------

    numeric_columns = [
        "Total Capacity (GB)",
        "Used Capacity (GB)",
        "Free Capacity (GB)",
        "Used Capacity (%)",
        "Free Capacity (%)"
    ]

    for column in numeric_columns:

        df[column] = pd.to_numeric(
            df[column],
            errors="coerce"
        )

    # --------------------------------------------------------
    # Remove invalid clusters
    # --------------------------------------------------------

    df = df[
        (df["Cluster"] != "")
        &
        (df["Cluster"].str.lower() != "nan")
    ]

    # --------------------------------------------------------
    # Fix Region values if required
    # --------------------------------------------------------

    try:

        from cluster_mapping import CLUSTER_MAPPING

        for index in df.index:

            cluster = df.at[
                index,
                "Cluster"
            ]

            region = str(
                df.at[
                    index,
                    "Region"
                ]
            ).strip()

            if region == "" or region.lower() == "nan":

                mapping = CLUSTER_MAPPING.get(
                    cluster
                )

                if mapping:

                    df.at[
                        index,
                        "Region"
                    ] = mapping.get(
                        "Region",
                        ""
                    )

                    if (
                        str(
                            df.at[
                                index,
                                "Country"
                            ]
                        ).strip()
                        == ""
                    ):

                        df.at[
                            index,
                            "Country"
                        ] = mapping.get(
                            "Country",
                            ""
                        )

                    if (
                        str(
                            df.at[
                                index,
                                "City"
                            ]
                        ).strip()
                        == ""
                    ):

                        df.at[
                            index,
                            "City"
                        ] = mapping.get(
                            "City",
                            ""
                        )

    except ImportError:

        pass

    # --------------------------------------------------------
    # Remove duplicate cluster/day records
    # --------------------------------------------------------

    df = (
        df
        .sort_values(
            [
                "Date",
                "Cluster"
            ]
        )
        .drop_duplicates(
            subset=[
                "Date",
                "Cluster"
            ],
            keep="last"
        )
        .reset_index(
            drop=True
        )
    )

    print(
        "Historical rows loaded : "
        + str(len(df))
    )

    print(
        "Unique dates           : "
        + str(df["Date"].nunique())
    )

    print(
        "Unique clusters        : "
        + str(df["Cluster"].nunique())
    )

    print()

    print("Regions:")

    print(
        df["Region"]
        .value_counts(
            dropna=False
        )
    )

    return df


# ============================================================
# BUILD WEEKLY SNAPSHOTS
# ============================================================

def get_weekly_snapshots(df):

    if df.empty:

        return (
            pd.DataFrame(),
            []
        )

    working = df.copy()

    # --------------------------------------------------------
    # Create week identifier.
    #
    # W-SUN means the week ends on Sunday.
    #
    # Example:
    # 21-Aug-2026 belongs to the week ending
    # 23-Aug-2026.
    # --------------------------------------------------------

    working["Week"] = (
        working["Date"]
        .dt.to_period("W-SUN")
    )

    available_weeks = sorted(
        working["Week"].unique()
    )

    selected_weeks = available_weeks[
        -NUMBER_OF_WEEKS:
    ]

    records = []

    # --------------------------------------------------------
    # Process each week
    # --------------------------------------------------------

    for week in selected_weeks:

        week_data = working[
            working["Week"] == week
        ].copy()

        if week_data.empty:
            continue

        # ----------------------------------------------------
        # Select Wednesday snapshot for this week.
        # Wednesday = dayofweek 2 in pandas.
        # The weekly report must use Wednesday data only.
        # ----------------------------------------------------

        wednesday_data = week_data[
            week_data["Date"].dt.dayofweek == 2
        ].copy()

        # If Wednesday data is not available for this week,
        # do not substitute another day.
        if wednesday_data.empty:
            continue

        snapshot_date = (
            wednesday_data["Date"]
            .max()
        )

        # ----------------------------------------------------
        # Use only Wednesday data for this week.
        # If duplicate cluster records exist on Wednesday,
        # keep the latest record for that cluster.
        # ----------------------------------------------------

        week_data = (
            wednesday_data
            .sort_values("Date")
            .groupby(
                "Cluster",
                as_index=False
            )
            .tail(1)
        )

        week_data = week_data[
            [
                "Cluster",
                "Region",
                "Country",
                "City",
                "Used Capacity (%)"
            ]
        ].copy()

        week_data["Snapshot_Date"] = (
            snapshot_date
        )

        records.append(
            week_data
        )

    if not records:

        return (
            pd.DataFrame(),
            []
        )

    weekly_data = pd.concat(
        records,
        ignore_index=True
    )

    return (
        weekly_data,
        selected_weeks
    )


# ============================================================
# BUILD WEEKLY TABLE
# ============================================================

def build_weekly_table(
    weekly_data,
    selected_weeks
):

    if weekly_data.empty:

        return pd.DataFrame()

    weekly_data["City Location"] = (
        weekly_data["City"]
        .astype(str)
        .str.strip()
    )

    # --------------------------------------------------------
    # Create actual date labels
    # --------------------------------------------------------

    snapshot_dates = (
        weekly_data[
            [
                "Snapshot_Date"
            ]
        ]
        .drop_duplicates()
        .sort_values(
            "Snapshot_Date"
        )
        .reset_index(
            drop=True
        )
    )

    date_labels = []

    for value in snapshot_dates[
        "Snapshot_Date"
    ]:

        date_labels.append(
            pd.Timestamp(
                value
            ).strftime(
                "%-d-%b-%y"
            )
        )

    # --------------------------------------------------------
    # Pivot by ACTUAL snapshot date
    # --------------------------------------------------------

    pivot = weekly_data.pivot_table(
        index=[
            "Cluster",
            "Region",
            "Country",
            "City Location"
        ],
        columns="Snapshot_Date",
        values="Used Capacity (%)",
        aggfunc="last"
    ).reset_index()

    pivot.columns.name = None

    # --------------------------------------------------------
    # Rename date columns
    # --------------------------------------------------------

    rename_map = {}

    for value in pivot.columns:

        if isinstance(
            value,
            pd.Timestamp
        ):

            rename_map[value] = (
                value.strftime(
                    "%-d-%b-%y"
                )
            )

    pivot = pivot.rename(
        columns=rename_map
    )

    # --------------------------------------------------------
    # Determine available week columns
    # --------------------------------------------------------

    week_columns = []

    for value in snapshot_dates[
        "Snapshot_Date"
    ]:

        label = pd.Timestamp(
            value
        ).strftime(
            "%-d-%b-%y"
        )

        if label in pivot.columns:

            week_columns.append(
                label
            )

    result_columns = [
        "Cluster",
        "Region",
        "Country",
        "City Location"
    ]

    result_columns.extend(
        week_columns
    )

    pivot = pivot[
        result_columns
    ]

    # --------------------------------------------------------
    # Sort by city
    # --------------------------------------------------------

    pivot = (
        pivot
        .sort_values(
            [
                "City Location",
                "Cluster"
            ]
        )
        .reset_index(
            drop=True
        )
    )

    return pivot


# ============================================================
# CREATE REGIONAL SHEET
# ============================================================

def create_region_sheet(
    wb,
    region_data,
    region_name,
    week_columns
):

    ws = wb.create_sheet(
        title=region_name
    )

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    ws["A1"] = (
        "Rubrik Weekly Capacity Growth"
    )

    ws["A1"].font = TITLE_FONT

    # --------------------------------------------------------
    # No Filters
    # --------------------------------------------------------


    # --------------------------------------------------------
    # Main table
    # --------------------------------------------------------

    start_row = 5

    headers = [
        "City Location",
        "TH %-Low",
        "TH %-High"
    ]

    headers.extend(
        week_columns
    )

    for column_number, header in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row=start_row,
            column=column_number,
            value=header
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER

    # --------------------------------------------------------
    # Write data
    # --------------------------------------------------------

    data_start_row = start_row + 1

    current_row = data_start_row

    if region_data.empty:

        # Still create an empty structure.
        current_row += 1

    else:

        for _, record in region_data.iterrows():

            city = record[
                "City Location"
            ]

            ws.cell(
                row=current_row,
                column=1,
                value=city
            )

            ws.cell(
                row=current_row,
                column=2,
                value=TH_LOW
            )

            ws.cell(
                row=current_row,
                column=3,
                value=TH_HIGH
            )

            for cell_column in [
                1,
                2,
                3
            ]:

                cell = ws.cell(
                    row=current_row,
                    column=cell_column
                )

                cell.border = BORDER

                if cell_column == 1:

                    cell.alignment = LEFT

                else:

                    cell.alignment = CENTER
                    cell.number_format = "0%"

            # ------------------------------------------------
            # Weekly values
            # ------------------------------------------------

            for index, week_column in enumerate(
                week_columns,
                start=4
            ):

                value = record.get(
                    week_column
                )

                if (
                    pd.isna(value)
                    or value == ""
                ):

                    value = None

                else:

                    value = float(
                        value
                    )

                cell = ws.cell(
                    row=current_row,
                    column=index,
                    value=value
                )

                cell.border = BORDER
                cell.alignment = CENTER
                cell.number_format = "0%"

            current_row += 1

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12

    for column_number in range(
        4,
        len(headers) + 1
    ):

        ws.column_dimensions[
            ws.cell(
                row=1,
                column=column_number
            ).column_letter
        ].width = 12

    ws.freeze_panes = "A6"

    # --------------------------------------------------------
    # RIGHT-SIDE CHART SOURCE TABLE
    # --------------------------------------------------------

    right_col = 10

    for index, header in enumerate(
        headers,
        start=right_col
    ):

        cell = ws.cell(
            row=start_row,
            column=index,
            value=header
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER

    right_data_start = (
        start_row + 1
    )

    right_row = right_data_start

    if not region_data.empty:

        for _, record in region_data.iterrows():

            ws.cell(
                row=right_row,
                column=right_col,
                value=record[
                    "City Location"
                ]
            )

            ws.cell(
                row=right_row,
                column=right_col + 1,
                value=TH_LOW
            )

            ws.cell(
                row=right_row,
                column=right_col + 2,
                value=TH_HIGH
            )

            for cell_offset in [
                1,
                2
            ]:

                cell = ws.cell(
                    row=right_row,
                    column=right_col + cell_offset
                )

                cell.number_format = "0%"
                cell.border = BORDER

            for offset, week_column in enumerate(
                week_columns,
                start=3
            ):

                value = record.get(
                    week_column
                )

                if (
                    pd.isna(value)
                    or value == ""
                ):

                    value = None

                else:

                    value = float(
                        value
                    )

                cell = ws.cell(
                    row=right_row,
                    column=right_col + offset,
                    value=value
                )

                cell.number_format = "0%"
                cell.border = BORDER

            right_row += 1

    # --------------------------------------------------------
    # CHART
    # --------------------------------------------------------

    if (
        not region_data.empty
        and len(week_columns) > 0
    ):

        first_data_row = (
            right_data_start
        )

        last_data_row = (
            right_row - 1
        )

        # -----------------------------------------------
        # Column chart
        # -----------------------------------------------

        bar_chart = BarChart()

        bar_chart.type = "col"

        bar_chart.style = 10

        bar_chart.title = ""

        bar_chart.y_axis.title = ""

        bar_chart.x_axis.title = ""

        bar_chart.y_axis.scaling.min = 0
        bar_chart.y_axis.scaling.max = 1

        bar_chart.y_axis.majorUnit = 0.25

        bar_chart.y_axis.numFmt = "0%"

        bar_chart.height = 10
        bar_chart.width = 18

        bar_chart.legend.position = "t"

        # -----------------------------------------------
        # Actual weekly capacity series
        # -----------------------------------------------

        for index in range(
            len(week_columns)
        ):

            data_col = (
                right_col + 3 + index
            )

            data = Reference(
                ws,
                min_col=data_col,
                min_row=start_row,
                max_row=last_data_row
            )

            categories = Reference(
                ws,
                min_col=right_col,
                min_row=first_data_row,
                max_row=last_data_row
            )

            bar_chart.add_data(
                data,
                titles_from_data=True
            )

            bar_chart.set_categories(
                categories
            )

        # -----------------------------------------------
        # Data labels
        # -----------------------------------------------

        bar_chart.dLbls = DataLabelList()

        bar_chart.dLbls.showVal = True

        bar_chart.dLbls.numFmt = "0%"

        bar_chart.dLbls.position = "outEnd"


        # -----------------------------------------------
        # Threshold line chart
        # -----------------------------------------------

        line_chart = LineChart()

        line_chart.y_axis.axId = 200

        line_chart.y_axis.crosses = "max"

        line_chart.y_axis.scaling.min = 0
        line_chart.y_axis.scaling.max = 1

        line_chart.y_axis.majorUnit = 0.25

        line_chart.y_axis.numFmt = "0%"

        line_chart.height = 10
        line_chart.width = 18

        # -----------------------------------------------
        # Low threshold
        # -----------------------------------------------

        low_data = Reference(
            ws,
            min_col=right_col + 1,
            min_row=start_row,
            max_row=last_data_row
        )

        line_chart.add_data(
            low_data,
            titles_from_data=True
        )

        # -----------------------------------------------
        # High threshold
        # -----------------------------------------------

        high_data = Reference(
            ws,
            min_col=right_col + 2,
            min_row=start_row,
            max_row=last_data_row
        )

        line_chart.add_data(
            high_data,
            titles_from_data=True
        )

        line_chart.set_categories(
            categories
        )

        # -----------------------------------------------
        # Combine
        # -----------------------------------------------

        bar_chart += line_chart

        # -----------------------------------------------
        # Put chart below tables
        # -----------------------------------------------

        ws.add_chart(
            bar_chart,
            "E14"
        )

    return ws


# ============================================================
# CREATE RAW DATA SHEET
# ============================================================

def create_raw_data_sheet(
    wb,
    df
):

    ws = wb.create_sheet(
        title="Raw Data"
    )

    # --------------------------------------------------------
    # Raw data columns
    # --------------------------------------------------------

    columns = [
        "Date",
        "Month",
        "Active",
        "Location",
        "Type",
        "Cluster",
        "Region",
        "Country",
        "City",
        "Total Capacity (GB)",
        "Used Capacity (GB)",
        "Free Capacity (GB)",
        "Used Capacity (%)",
        "Free Capacity (%)",
        "TH High",
        "TH Low"
    ]

    for column_number, column_name in enumerate(
        columns,
        start=1
    ):

        cell = ws.cell(
            row=1,
            column=column_number,
            value=column_name
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER

    # --------------------------------------------------------
    # Write history data
    # --------------------------------------------------------

    current_row = 2

    for _, record in df.sort_values(
        [
            "Date",
            "Cluster"
        ]
    ).iterrows():

        values = [
            record["Date"],
            record["Date"].strftime("%Y-%m"),
            record["Active"],
            record["Location"],
            record["Type"],
            record["Cluster"],
            record["Region"],
            record["Country"],
            record["City"],
            record["Total Capacity (GB)"],
            record["Used Capacity (GB)"],
            record["Free Capacity (GB)"],
            record["Used Capacity (%)"],
            record["Free Capacity (%)"],
            TH_HIGH,
            TH_LOW
        ]

        for column_number, value in enumerate(
            values,
            start=1
        ):

            cell = ws.cell(
                row=current_row,
                column=column_number,
                value=value
            )

            cell.border = BORDER
            if column_number == 1:
                cell.number_format = "yyyy-mm-dd"


            elif column_number in [
                10,
                11,
                12
            ]:

                cell.number_format = "#,##0.00"

            elif column_number in [
                13,
                14,
                15,
                16
            ]:

                cell.number_format = "0%"

        current_row += 1

    # --------------------------------------------------------
    # Widths
    # --------------------------------------------------------

    widths = {
        "A": 12,
        "B": 12,
        "C": 38,
        "D": 24,
        "E": 16,
        "F": 12,
        "G": 16,
        "H": 20,
        "I": 18,
        "J": 18,
        "K": 18,
        "L": 18,
        "M": 18,
        "N": 12,
        "O": 12
    }

    for column, width in widths.items():

        ws.column_dimensions[
            column
        ].width = width

    ws.freeze_panes = "A2"

    return ws


# ============================================================
# GENERATE WEEKLY REPORT
# ============================================================

def generate_weekly_report():

    print()
    print("=" * 70)
    print("Rubrik Weekly Capacity Report")
    print("=" * 70)

    # --------------------------------------------------------
    # Load history
    # --------------------------------------------------------

    history = load_history()

    # --------------------------------------------------------
    # Build weekly snapshots
    # --------------------------------------------------------

    weekly_data, selected_weeks = (
        get_weekly_snapshots(
            history
        )
    )

    if weekly_data.empty:

        raise ValueError(
            "No weekly capacity data available."
        )

    # --------------------------------------------------------
    # Show selected actual snapshot dates
    # --------------------------------------------------------

    selected_dates = (
        weekly_data[
            "Snapshot_Date"
        ]
        .drop_duplicates()
        .sort_values()
        .tolist()
    )

    print()

    print(
        "Weeks included : "
        + str(len(selected_dates))
    )

    for snapshot_date in selected_dates:

        print(
            "  "
            + pd.Timestamp(
                snapshot_date
            ).strftime(
                "%d %b %Y"
            )
        )

    # --------------------------------------------------------
    # Build weekly table
    # --------------------------------------------------------

    weekly_table = build_weekly_table(
        weekly_data,
        selected_weeks
    )

    if weekly_table.empty:

        raise ValueError(
            "Weekly table is empty."
        )

    # --------------------------------------------------------
    # Get week columns
    # --------------------------------------------------------

    week_columns = []

    for column in weekly_table.columns:

        if column in [
            "Cluster",
            "Region",
            "Country",
            "City Location"
        ]:

            continue

        week_columns.append(
            column
        )

    print()

    print(
        "Clusters included : "
        + str(len(weekly_table))
    )

    print(
        "Weekly columns     : "
        + ", ".join(
            week_columns
        )
    )

    # --------------------------------------------------------
    # Create output directory
    # --------------------------------------------------------

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Output filename
    # --------------------------------------------------------

    today = datetime.now()

    report_name = today.strftime(
        "Backup_Capacity_Weekly_%d_%b_%Y.xlsx"
    )

    report_path = os.path.join(
        OUTPUT_DIR,
        report_name
    )

    # --------------------------------------------------------
    # Create workbook
    # --------------------------------------------------------

    wb = Workbook()

    default_sheet = wb.active

    wb.remove(
        default_sheet
    )

    # --------------------------------------------------------
    # Regions
    # --------------------------------------------------------

    regions = [
        "NA",
        "NA_DD",
        "LATAM",
        "EU",
        "APAC",
        "CHINA"
    ]

    for region in regions:

        if region == "CHINA":

            region_data = weekly_table[
                weekly_table[
                    "Country"
                ]
                .astype(str)
                .str.upper()
                .eq("CHINA")
            ].copy()

        elif region == "NA_DD":

            region_data = weekly_table.iloc[
                0:0
            ].copy()

        else:

            region_data = weekly_table[
                weekly_table[
                    "Region"
                ]
                .astype(str)
                .str.upper()
                .eq(region)
            ].copy()

        create_region_sheet(
            wb,
            region_data,
            region,
            week_columns
        )

    # --------------------------------------------------------
    # Raw Data
    # --------------------------------------------------------

    create_raw_data_sheet(
        wb,
        history
    )

    # --------------------------------------------------------
    # Sheet order
    # --------------------------------------------------------

    desired_order = [
        "NA",
        "NA_DD",
        "LATAM",
        "EU",
        "APAC",
        "CHINA",
        "Raw Data"
    ]

    wb._sheets = [
        wb[name]
        for name in desired_order
        if name in wb.sheetnames
    ]

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    wb.save(
        report_path
    )

    file_exists = os.path.exists(
        report_path
    )

    print()
    print("=" * 70)
    print("Weekly Excel Report Generated")
    print("=" * 70)

    print(
        "Report Location : "
        + report_path
    )

    print(
        "Absolute Path   : "
        + os.path.abspath(
            report_path
        )
    )

    print(
        "Clusters        : "
        + str(len(weekly_table))
    )

    print(
        "Weeks Included  : "
        + str(len(selected_dates))
    )

    print(
        "File Exists     : "
        + str(file_exists)
    )

    print(
        "Sheets          : "
        + ", ".join(
            wb.sheetnames
        )
    )

    print("=" * 70)

    return report_path


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    report_file = (
        generate_weekly_report()
    )

    print()

    print(
        "Sending Weekly Capacity Report Email..."
    )

    try:

        from email_report import send_email

        send_email(
            report_file
        )

    except Exception as exc:

        print(
            "Email sending failed: "
            + str(exc)
        )
